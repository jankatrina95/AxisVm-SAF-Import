"""Generischer Leser fuer SAF-Dateien (.xlsx).

Jedes Arbeitsblatt wird als Liste von Zeilen eingelesen; die erste Zeile
liefert die Spaltenkoepfe. Leere Zeilen werden uebersprungen.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

Row = dict[str, Any]


@dataclass
class SafModel:
    """Inhalt einer SAF-Datei: Blattname -> Zeilen."""

    source: Path
    sheets: dict[str, list[Row]] = field(default_factory=dict)

    def rows(self, sheet: str) -> list[Row]:
        """Zeilen eines Blatts; leere Liste, wenn das Blatt fehlt."""
        return self.sheets.get(sheet, [])

    def summary(self) -> dict[str, int]:
        """Zeilenanzahl je Blatt."""
        return {name: len(rows) for name, rows in self.sheets.items()}


def read_saf(path: str | Path) -> SafModel:
    """Liest eine SAF-.xlsx-Datei vollstaendig ein."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"SAF-Datei nicht gefunden: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        model = SafModel(source=path)
        for ws in wb.worksheets:
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if not header:
                model.sheets[ws.title] = []
                continue
            keys = [str(h).strip() if h is not None else "" for h in header]
            rows: list[Row] = []
            for values in rows_iter:
                if values is None or all(v is None for v in values):
                    continue
                rows.append({k: v for k, v in zip(keys, values) if k})
            model.sheets[ws.title] = rows
        return model
    finally:
        wb.close()
